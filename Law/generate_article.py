import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_document():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(9)

    for sname in ["Heading 1", "Heading 2", "Heading 3"]:
        s = styles[sname]
        s.font.name = "Times New Roman"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        s.font.size = Pt(14)
        s.font.bold = True
        s.paragraph_format.line_spacing = 1.5
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(8)
        s.paragraph_format.keep_with_next = True

    def para(text="", bold=False, align=None, after=9):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)
        r.bold = bold
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(after)
        return p

    def h(text, level=1):
        p = doc.add_paragraph(style=f"Heading {level}")
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(14)
        r.bold = True
        return p

    def numbered(text):
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(12)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("1. COMMUNITY SERVICE AS A PUNISHMENT UNDER THE BHARATIYA NYAYA SANHITA, 2023")
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(14)
    r.bold = True

    # 2
    h("2. INTRODUCTION")
    paras = [
        "Punishment is one of the most important functions of the criminal justice system. Traditionally, criminal law has relied upon imprisonment, fine and, in the gravest cases, life imprisonment or the death penalty. However, modern criminal jurisprudence recognises that punishment should not be confined to inflicting suffering upon an offender. It should also promote rehabilitation, accountability, social responsibility and, where possible, restoration of the harm caused to society.",
        "The enactment of the Bharatiya Nyaya Sanhita, 2023 (BNS) marks an important development in Indian criminal law by expressly recognising community service as a form of punishment. Section 4(f) of the BNS places community service among the punishments available under the Sanhita. This is significant because community service was not expressly recognised as a substantive category of punishment under the Indian Penal Code, 1860.",
        "Community service represents a movement away from an exclusively punitive model towards a more reformative and restorative approach. Instead of sending every minor offender to prison, the law permits the offender, in specified circumstances, to perform useful work for the benefit of the community without remuneration. The objective is not to excuse criminal conduct but to make the offender confront the consequences of his conduct in a constructive manner.",
        "The Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS) further provides the operational definition of community service. Section 23 explains it as work which a Court may order a convict to perform as a form of punishment that benefits the community, without entitlement to remuneration.",
        "The introduction of community service therefore raises important legal questions: What is the scope of community service under the BNS? For which offences can it be imposed? How should courts determine its duration and nature? Does it adequately balance punishment, deterrence and rehabilitation? What safeguards are necessary to prevent arbitrary or unequal sentencing?",
        "This article examines these questions and argues that community service has the potential to become an effective component of Indian criminal justice, provided that it is implemented through clear guidelines, judicially controlled discretion, adequate supervision and respect for proportionality and constitutional values."
    ]
    for x in paras: para(x)

    # 3
    h("3. LEGAL FRAMEWORK")
    h("3.1 Section 4 of the Bharatiya Nyaya Sanhita, 2023", 2)
    para("Section 4 of the BNS deals with punishments. It recognises the following forms of punishment:")
    for x in ["Death", "Imprisonment for life", "Imprisonment, rigorous or simple", "Forfeiture of property", "Fine", "Community service"]:
        numbered(x)
    para("The inclusion of community service as clause (f) is a major legislative change. It demonstrates that imprisonment is not necessarily the only appropriate response to every criminal offence.")
    para("However, community service is not a universal punishment available for every offence. The BNS specifically attaches community service to certain offences. The statutory categories discussed in this article are Sections 202, 209, 226, 303(2), 355 and 356(2). Thus, the legislature has adopted a controlled rather than unlimited model of community sentencing.")

    h("3.2 Meaning of Community Service under Section 23 of BNSS", 2)
    for x in [
        "Section 23 of the BNSS is particularly important because it provides the statutory meaning of community service.",
        "It states, in substance, that community service means work which the Court orders a convict to perform as punishment for the benefit of the community and for which the convict receives no remuneration.",
        "The definition contains three essential elements. First, the work must be ordered by a Court. Second, the work must benefit the community. Third, the offender cannot claim payment for performing the service.",
        "Therefore, community service is not ordinary voluntary social work. It is a court-imposed criminal sentence. Non-compliance cannot simply be treated as withdrawal from voluntary activity because the service forms part of the punishment imposed following conviction."
    ]:
        para(x)

    h("3.3 Powers of Magistrates", 2)
    para("Section 23 of the BNSS also determines the sentencing powers of Magistrates. A Magistrate of the First Class may impose community service, while a Magistrate of the Second Class may also impose community service within the framework of offences legally permitting that punishment. This is particularly important because most offences for which community service is prescribed are comparatively minor and are likely to be dealt with by Magistrates. Consequently, the practical success of community sentencing will depend substantially upon the quality of judicial discretion exercised by trial courts.")

    h("3.4 Offences for Which Community Service Is Available", 2)
    off = [
        ("3.4.1 Section 202 – Public Servant Unlawfully Engaging in Trade",
         "Section 202 deals with a public servant unlawfully engaging in trade. The availability of community service in such a case reflects the principle that not every violation by a public servant necessarily requires imprisonment. Depending upon the circumstances and seriousness of the misconduct, a court may select a non-custodial punishment. Any assignment should remain socially useful, authorised under applicable rules and proportionate to the offence."),
        ("3.4.2 Section 209 – Non-Appearance in Response to Proclamation",
         "Section 209 deals with failure to appear at the specified place and time pursuant to a proclamation under Section 84 of the BNSS. The provision permits imprisonment, fine, or community service for the basic offence. However, where the person has been declared a proclaimed offender under the relevant provision, the punishment becomes significantly more serious. This distinction shows that community service is not intended to replace serious punishment where conduct demonstrates aggravated disregard for the criminal process."),
        ("3.4.3 Section 226 – Attempt to Commit Suicide to Compel or Restrain Exercise of Lawful Power",
         "Section 226 concerns an attempt to commit suicide for the purpose of compelling or restraining the exercise of lawful power by a public servant. Where the offence falls within the statutory requirements, community service may be considered as one of the available punishments. The provision also raises an important jurisprudential issue: punishment must remain proportionate to culpability and must account for the precise circumstances in which the offence was committed."),
        ("3.4.4 Section 303(2) – First-Time Petty Theft",
         "One of the most significant applications of community service is found in Section 303(2). Where the value of stolen property is less than ₹5,000, the offender is convicted for the first time, and the stolen property or its value is returned or restored, the statutory framework provides for community service. This provision demonstrates the reformative philosophy of the BNS particularly clearly. Imprisonment for a low-value first offence could disrupt employment, education and family relationships, whereas community service can impose a real consequence while allowing the offender to remain connected with society. The provision therefore attempts to balance accountability, restitution, rehabilitation and social protection."),
        ("3.4.5 Section 355 – Misconduct in Public by a Drunken Person",
         "Section 355 provides punishment for a person who, while intoxicated, appears in a public place or a place where he is a trespasser and conducts himself in a manner causing annoyance to another person. The punishment may include simple imprisonment up to twenty-four hours, fine up to ₹1,000, or both, or community service. This offence provides a particularly appropriate field for community sentencing because the harm involved is generally public nuisance or annoyance rather than serious physical injury. Community service can convert the offender's wrongdoing into a constructive contribution to the same community that was disturbed."),
        ("3.4.6 Section 356(2) – Defamation",
         "Section 356 deals with defamation. Under Section 356(2), a person who defames another may be punished with simple imprisonment up to two years, fine, both, or community service. The inclusion of community service in defamation cases is legally interesting because defamation involves a conflict between reputation, individual dignity and freedom of expression. The court must carefully determine whether community service is proportionate to the circumstances, considering the nature of the statement, intention, harm caused, subsequent conduct and the statutory framework governing defamation.")
    ]
    for a, b in off:
        h(a, 3); para(b)

    # 4
    h("4. CASE LAWS")
    cases = [
        ("4.1 State of Punjab v. Prem Sagar (2008)",
         "In State of Punjab v. Prem Sagar, (2008) 7 SCC 550, the Supreme Court highlighted the difficulties surrounding sentencing discretion in India. The Court recognised that Indian sentencing law has historically lacked comprehensive sentencing guidelines and emphasised the need for judicial discretion to be exercised objectively, taking into account the circumstances of the offence and offender. The decision is highly relevant to community service because introducing a new punishment gives trial courts greater responsibility in deciding the appropriate sentence. Community service cannot be imposed mechanically; the court must consider whether it is proportionate to the offence and appropriate for the particular offender."),
        ("4.2 Soman v. State of Kerala (2013)",
         "In Soman v. State of Kerala, (2013) 11 SCC 382, the Supreme Court explained important sentencing principles. The Court recognised proportionality and deterrence as major sentencing rationales and stated that the sentence should correspond to the seriousness or gravity of the offence. This principle provides a strong judicial foundation for community service. Community service should not be understood as a lenient sentence merely because it does not involve imprisonment. A properly designed community-service order can constitute a meaningful deprivation of time and personal convenience while simultaneously producing social benefit."),
        ("4.3 Manoj Kumar v. State (Government of NCT of Delhi)",
         "Before the BNS expressly recognised community service as a statutory punishment, Indian High Courts had occasionally used community-oriented directions while granting relief in appropriate cases. In Manoj Kumar v. State (Government of NCT of Delhi), the Delhi High Court directed community service at Lok Nayak Jai Prakash Narayan Hospital in an appropriate criminal proceeding. The case illustrates that Indian courts had already begun experimenting with constructive forms of justice even before community service became expressly recognised under the BNS. The BNS now gives the concept a statutory foundation in specified offences."),
        ("4.4 Solemen SK v. State of West Bengal",
         "In Solemen SK v. State of West Bengal, the Supreme Court directed an individual to plant trees as a form of community-oriented service in the context of the relief granted by the Court. Although this case did not arise under Section 4(f) of the BNS, it demonstrates the Indian judiciary's recognition of constructive and socially beneficial measures as an appropriate component of judicial relief in suitable circumstances.")
    ]
    for a, b in cases:
        h(a, 2); para(b)

    # 5
    h("5. ANALYSIS")
    analysis = [
        ("5.1 From Retributive Justice to Reformative Justice",
         "Traditional punishment is often associated with retribution: an offender commits a wrong and therefore deserves punishment. Modern criminal justice, however, recognises several objectives, including retribution, deterrence, prevention, rehabilitation and restoration. Community service attempts to combine these objectives. The offender is punished because the work is compulsory and unpaid. At the same time, the offender contributes something useful to society. This can encourage responsibility rather than merely producing suffering. Therefore, community service represents a movement towards constructive punishment."),
        ("5.2 Community Service and Prison Overcrowding",
         "India has long faced pressure on its prison infrastructure. Short-term imprisonment for minor offences can be particularly problematic. A person imprisoned for a minor offence may lose employment, become separated from family and potentially come into contact with more serious offenders. For first-time and low-risk offenders, community service can therefore act as an alternative that preserves social bonds while still recognising criminal responsibility. However, community service should not be used merely to reduce prison population statistics. The central question must always be whether it is a just and proportionate sentence in the particular case."),
        ("5.3 Rehabilitation of the Offender",
         "The strongest justification for community service is rehabilitation. A person who commits a minor offence may not necessarily be a hardened criminal. Imprisonment can sometimes increase rather than reduce criminal tendencies. Community service can provide an opportunity for the offender to recognise the consequences of wrongdoing, develop discipline, understand social responsibility, contribute to the community and avoid unnecessary exposure to prison. The punishment can therefore communicate that wrongdoing has consequences while still recognising the possibility of reform."),
        ("5.4 Restorative Justice",
         "Community service is closely connected with the philosophy of restorative justice. Retributive justice primarily asks, “What punishment does the offender deserve?” Restorative justice additionally asks, “How can the harm caused by wrongdoing be addressed and how can the offender be reintegrated into society?” Community service does not always directly compensate the victim. Nevertheless, it can restore a broader social balance by requiring the offender to make a constructive contribution. It should, however, operate alongside any lawful victim compensation or restitution rather than automatically replacing it."),
        ("5.5 Proportionality",
         "The principle of proportionality is central. A community-service order should correspond to the seriousness of the offence, harm caused, intention of the offender, previous criminal record, age and circumstances of the offender, possibility of rehabilitation and impact of the punishment. The Supreme Court's sentencing jurisprudence repeatedly emphasises proportionality, deterrence and rehabilitation. Thus, judicial discretion should not become arbitrary discretion. The reasons for selecting community service and determining its duration should be apparent from the sentencing decision."),
        ("5.6 Equality Before Law",
         "Article 14 of the Constitution guarantees equality before law and equal protection of laws. This principle has relevance to community sentencing because similarly situated offenders should ordinarily receive reasonably consistent treatment. If one court orders a short period of service for an offence while another orders a substantially longer period for materially similar conduct without adequate reasoning, sentencing disparity may arise. Clear sentencing guidelines can therefore promote consistency while preserving judicial discretion for exceptional circumstances."),
        ("5.7 Article 21 and Dignity",
         "Article 21 protects life and personal liberty and has been interpreted by the Supreme Court as encompassing dignity and fairness. Community service can be consistent with Article 21 if it is authorised by law, imposed by a competent court, proportionate, reasonable, non-discriminatory and implemented with dignity. The work should never become humiliating, degrading or exploitative. A punishment intended to reform an offender should not itself violate the offender's basic human dignity.")
    ]
    for a, b in analysis:
        h(a, 2); para(b)

    # 6
    h("6. RECENT DEVELOPMENTS")
    recent = [
        ("6.1 Maharashtra Community Service Rules, 2025",
         "One of the most significant developments has occurred in Maharashtra. The Maharashtra (Community Services as Punishment for Certain Offences) Rules, 2025 were notified on 28 November 2025. The Rules were framed under Section 23 of the BNSS read with Section 4(f) of the BNS. The Rules provide greater clarity concerning the nature of community service, place of service, supervising authority, monitoring authority and duration. They provide that community service may ordinarily range from one to thirty-one days or forty to 240 hours, depending upon the applicable framework, and service is limited to eight hours per day. The Schedule includes activities such as cleaning and maintenance of government hospitals, legal aid-related assistance and other public-oriented duties. This is an important step because statutory recognition must be supported by an administrative mechanism capable of supervising and verifying completion."),
        ("6.2 Pune: Practical Implementation",
         "In December 2025, a Pune court reportedly imposed four days of community service on two persons convicted under Section 355 for drunken misconduct in a public place. The accused were ordered to perform three hours of service per day, with duties involving cleanliness of police-station premises and assisting traffic police in public-awareness activities. This example demonstrates how community service can be tailored to the nature of the offence. Rather than simply imposing a fine, the court linked the punishment to a socially useful activity."),
        ("6.3 Jammu and Kashmir Courts",
         "In early 2025, courts in Jammu also imposed community-service sentences for public misconduct by intoxicated persons. In one reported case, the offender was directed to undertake cleaning and maintenance work at a public health facility. These cases represent the early practical development of community sentencing under the BNS and show the importance of having institutions capable of receiving and supervising persons sentenced to community service."),
        ("6.4 Karnataka: Tree Plantation as Community Service",
         "A recent development in Karnataka demonstrates an increasingly creative approach. In August 2026, a Belagavi court ordered a person convicted of public misconduct while intoxicated to plant 50 saplings, divided between the court premises and a police station, along with a fine. The order illustrates the possibility of connecting community service with environmental and civic objectives. However, such innovative sentencing must remain within statutory authority and applicable State rules. Judicial creativity cannot replace legislative standards."),
        ("6.5 Expansion of State-Level Frameworks",
         "The implementation of community service has gradually moved from statutory recognition to administrative regulation. Recent legal commentary notes that Maharashtra, West Bengal, Assam, Manipur, Haryana and Uttar Pradesh have developed frameworks or guidelines concerning community service, including matters such as duration, nature of work, locations and supervision. This development is significant because a punishment becomes meaningful only when there is a reliable mechanism for implementation and verification.")
    ]
    for a, b in recent:
        h(a, 2); para(b)

    # 7
    h("7. ORIGINALITY")
    orig = [
        "This article is based on an independent analysis of the statutory framework governing community service under the Bharatiya Nyaya Sanhita, 2023 (BNS) and the Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS), together with relevant judicial decisions and recent developments in sentencing policy.",
        "The legal propositions discussed in this article are based primarily on statutory provisions, judicial decisions and government rules relating to community service. Particular attention has been given to Section 4(f) of the BNS, which recognises community service as a form of punishment, and Section 23 of the BNSS, which provides the statutory framework for community service.",
        "The article also examines relevant principles of Indian sentencing jurisprudence, including proportionality, deterrence, rehabilitation, judicial discretion, equality before law and protection of human dignity. Judicial decisions including State of Punjab v. Prem Sagar and Soman v. State of Kerala have been considered to understand the principles governing sentencing discretion.",
        "Recent developments, including the Maharashtra (Community Services as Punishment for Certain Offences) Rules, 2025, have also been examined to understand the practical implementation of community-service orders.",
        "The discussion has been independently organised and analysed rather than reproduced from any single secondary source. The article seeks to maintain academic originality while relying upon verified legal sources, including legislation, reported judicial decisions and government notifications. All statutory provisions, judicial authorities and recent developments referred to in the article should be verified against their authoritative sources before final submission."
    ]
    for x in orig: para(x)

    # 8
    h("8. CONCLUSION")
    concl = [
        "The introduction of community service under the Bharatiya Nyaya Sanhita, 2023 represents an important development in India's sentencing policy. By expressly recognising community service as a punishment, the legislature has expanded the range of sentencing options available to courts beyond imprisonment and monetary penalties.",
        "Section 4(f) of the BNS establishes community service as one of the punishments that may be imposed under the new criminal law framework, while the BNSS provides the procedural foundation for its application. The present statutory scheme, however, does not make community service a general substitute for imprisonment. Its application is confined to specified offences and circumstances, particularly offences considered comparatively less serious. This indicates that the legislative approach seeks to introduce community-based punishment while retaining imprisonment for more serious criminal conduct.",
        "Community service also reflects an important change in the philosophy of punishment. Traditional criminal justice has frequently focused on imprisonment, deterrence and monetary penalties. Community service introduces an additional objective: requiring the offender to undertake constructive work for the benefit of society. It therefore has the potential to combine accountability, rehabilitation and social responsibility.",
        "At the same time, statutory recognition alone cannot guarantee effective implementation. Community service requires appropriate supervision, clear judicial directions, reliable mechanisms for recording completion and safeguards against arbitrary or degrading assignments. Courts must ensure that the punishment remains proportionate to the offence and does not become discriminatory or humiliating. The interests of victims must likewise be protected, since community service should not automatically replace compensation, restitution or any other remedy available under law.",
        "The principle of individualised sentencing is particularly important. Courts should consider the nature and circumstances of the offence, the offender's background, previous criminal history and prospects of rehabilitation before imposing community service. Clear guidelines can promote consistency while preserving the judicial discretion necessary to deal with individual cases.",
        "Recent developments such as the Maharashtra Community Services as Punishment for Certain Offences Rules, 2025 demonstrate the importance of developing an institutional mechanism alongside the statutory provision. Effective implementation will require coordination among courts, supervising authorities, local institutions and other agencies responsible for providing and monitoring community-service assignments.",
        "The clear legal takeaway is that community service should not be treated as merely a lenient alternative to imprisonment. Properly implemented, it represents a constructive form of punishment in which the offender remains accountable while being given an opportunity to contribute positively to society. Its success will ultimately depend upon proportionality, dignity, supervision, consistency and effective enforcement.",
        "Therefore, Indian criminal justice should focus not merely on the quantity of punishment imposed but also on whether the punishment achieves legitimate sentencing objectives. Community service, when applied to suitable offences and offenders within a clearly regulated framework, can contribute to rehabilitation, reduce unnecessary reliance on imprisonment and strengthen the social purpose of criminal justice.",
        "The future development of community sentencing should consequently focus on clear national standards, individualised judicial assessment, effective monitoring and protection of both offender and victim interests. If these safeguards are maintained, community service can develop into a credible and meaningful component of India's modern sentencing system."
    ]
    for x in concl: para(x)

    # Add footer with page number
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Page ")
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        p._p.append(fld)

    target_dir = r"C:\Users\adity\Desktop\Subjects\Law"
    os.makedirs(target_dir, exist_ok=True)
    out_path = os.path.join(target_dir, "Community_Service_BNS_Article_2000_3500_Polished.docx")
    doc.save(out_path)

    # approximate word count from document text
    word_count = sum(len(p.text.split()) for p in doc.paragraphs)
    print(f"SAVED: {out_path}")
    print(f"WORD_COUNT: {word_count}")

if __name__ == "__main__":
    create_document()
